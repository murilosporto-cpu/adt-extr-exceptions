from playwright.sync_api import sync_playwright
import time
import os
import openpyxl

def apply_and_refresh(page, iso_date):
    print(f"\n--- Aplicando {iso_date} ---", flush=True)
    page.keyboard.press("Escape")
    time.sleep(0.5)

    page.evaluate("document.querySelector('#date-selection').click()")
    time.sleep(1)

    page.evaluate("""() => {
        const item = Array.from(document.querySelectorAll('.dx-menu-item'))
            .find(el => el.textContent.trim() === 'Custom');
        if (item) item.click();
    }""")
    time.sleep(1.5)

    page.evaluate("""(dStr) => {
        const parts = dStr.split('-');
        const dObj = new Date(parseInt(parts[0]), parseInt(parts[1]) - 1, parseInt(parts[2]));
        
        $('#custom_date_begin_selector').dxDateBox('instance').option('value', dObj);
        $('#custom_date_end_selector').dxDateBox('instance').option('value', dObj);
        
        if (PwrJSDateFilters._selectedDateFilterOption) {
            PwrJSDateFilters._selectedDateFilterOption.BeginDate = dObj;
            PwrJSDateFilters._selectedDateFilterOption.EndDate = dObj;
        }
        
        PwrJSSpa.RefreshPWRReport();
    }""", iso_date)

    page.keyboard.press("Escape")

    time.sleep(2)
    page.wait_for_function("""() => {
        const inCb = typeof ASPxGridViewMainReport !== 'undefined' && ASPxGridViewMainReport.InCallback && ASPxGridViewMainReport.InCallback();
        const loadVisible = $('.dx-loadpanel:visible').length > 0;
        return !inCb && !loadVisible;
    }""", timeout=30000)
    time.sleep(2)

def export_file(page, filename):
    page.evaluate("document.querySelector('.dx-icon-menu').click()")
    time.sleep(1.5)
    with page.expect_download(timeout=30000) as dl_info:
        page.click("text='Export to Excel (All Columns)'")
    dl = dl_info.value
    dest = os.path.join(r'c:\Users\muril\OneDrive\FRANQUIAS\master mind\cafe-com-pwr\novo projeto', filename)
    dl.save_as(dest)
    print(f"Salvo: {filename} ({os.path.getsize(dest)} bytes)", flush=True)
    page.keyboard.press("Escape")
    time.sleep(1)
    return dest

with sync_playwright() as p:
    browser = p.chromium.launch(headless=True)
    context = browser.new_context(accept_downloads=True)
    page = context.new_page()
    page.goto('https://pwr.dominos.com')
    page.fill('#txtUsername', 'portom')
    page.fill('#txtPassword', '!!dominos@2026!!')
    with page.expect_navigation(timeout=45000):
        page.evaluate("""() => {
            const d = new Date();
            document.querySelector('#txtTZOffSet').value = d.getTimezoneOffset();
            __doPostBack('btnLogin', '');
        }""")
    time.sleep(4)

    # Seleciona All Stores (Stores)
    page.click("#scope-selection")
    time.sleep(2)
    page.locator(".dx-treeview-item:has-text('All Stores (Stores)')").first.click()
    time.sleep(6)

    # Navega para KEYS Service Exceptions
    print("Navegando para Service Exceptions...")
    page.click("//*[contains(text(), 'KEYS')]")
    time.sleep(2)
    page.evaluate("""() => {
        const items = Array.from(document.querySelectorAll('.dx-treeview-item'));
        const serviceItems = items.filter(el => el.textContent.trim() === 'Service');
        if (serviceItems.length >= 2) {
            serviceItems[1].click();
        } else if (serviceItems.length === 1) {
            serviceItems[0].click();
        }
    }""")
    time.sleep(3.5)
    page.click("text='Service Exceptions'")
    time.sleep(6)

    # Teste 1: 17 de Agosto
    apply_and_refresh(page, "2026-08-17")
    dest17 = export_file(page, "test_exc_17.xlsx")

    # Teste 2: 18 de Agosto
    apply_and_refresh(page, "2026-08-18")
    dest18 = export_file(page, "test_exc_18.xlsx")

    browser.close()

print("\n--- COMPARAÇÃO DOS DADOS DE SERVICE EXCEPTIONS ---")
for fn in [dest17, dest18]:
    wb = openpyxl.load_workbook(fn, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    cols = ['Store', 'Order Count', 'Singles %', 'Extreme %', 'eADT']
    idxs = [headers.index(c) for c in cols if c in headers]
    for r in range(2, 10):
        row_vals = [ws[r][i].value for i in idxs]
        if row_vals[1] is not None and row_vals[1] > 0:
            print(f"{os.path.basename(fn)} | Loja {row_vals[0]}: {row_vals}")
            break
