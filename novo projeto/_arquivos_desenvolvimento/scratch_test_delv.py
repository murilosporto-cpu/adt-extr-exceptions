
import openpyxl

days = [f'2026-08-{d:02d}' for d in range(17, 24)]
daily_stats = []

for d in days:
    f_sum = f'novo projeto/dados_all_stores/Keys Summary - All Stores (Stores) ({d}).xlsx'
    f_exc = f'novo projeto/dados_all_stores/KEYS Service Exceptions - All Stores (Stores) ({d}).xlsx'
    
    wb_s = openpyxl.load_workbook(f_sum, data_only=True)
    ws_s = wb_s.active
    h_s = [c for c in next(ws_s.iter_rows(values_only=True))]
    
    wb_e = openpyxl.load_workbook(f_exc, data_only=True)
    ws_e = wb_e.active
    h_e = [c for c in next(ws_e.iter_rows(values_only=True))]
    
    row_s = None
    for r in ws_s.iter_rows(values_only=True):
        if r[h_s.index('Store')] and str(r[h_s.index('Store')]) == '19762':
            row_s = r
            break
            
    row_e = None
    for r in ws_e.iter_rows(values_only=True):
        if r[h_e.index('Store')] and str(r[h_e.index('Store')]) == '19762':
            row_e = r
            break
            
    total_ord = row_s[h_s.index('Order Count')]
    eadt = row_s[h_s.index('eADT')]
    ext = row_s[h_s.index('% of Est Extreme Deliveries')]
    delv_ord = row_e[h_e.index('Delv Order Count')]
    
    daily_stats.append({
        'day': d,
        'tot_orders': total_ord,
        'delv_orders': delv_ord,
        'eadt': eadt,
        'ext': ext
    })

tot_delv = sum(x['delv_orders'] for x in daily_stats)
tot_ord = sum(x['tot_orders'] for x in daily_stats)

for x in daily_stats:
    print(x['day'], 'tot:', x['tot_orders'], 'delv:', x['delv_orders'], 'eadt:', round(x['eadt'], 2), 'ext:', round(x['ext']*100, 2))

print('Total Orders:', tot_ord)
print('Total Delv:', tot_delv)

eadt_by_tot = sum(x['eadt'] * x['tot_orders'] for x in daily_stats) / tot_ord
ext_by_tot = sum(x['ext'] * x['tot_orders'] for x in daily_stats) / tot_ord
print('Weighted by TOTAL orders -> eADT:', round(eadt_by_tot, 4), 'ext:', round(ext_by_tot*100, 4))

eadt_by_delv = sum(x['eadt'] * x['delv_orders'] for x in daily_stats) / tot_delv
ext_by_delv = sum(x['ext'] * x['delv_orders'] for x in daily_stats) / tot_delv
print('Weighted by DELV orders  -> eADT:', round(eadt_by_delv, 4), 'ext:', round(ext_by_delv*100, 4))
