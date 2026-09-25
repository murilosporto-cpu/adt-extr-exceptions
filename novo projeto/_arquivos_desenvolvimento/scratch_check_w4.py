
import openpyxl

days = [f'2026-09-{d:02d}' for d in range(7, 14)]
print('Checking days 07 to 13 for Rio Branco (19762):')

for d in days:
    f_sum = f'novo projeto/dados_all_stores/Keys Summary - All Stores (Stores) ({d}).xlsx'
    f_exc = f'novo projeto/dados_all_stores/KEYS Service Exceptions - All Stores (Stores) ({d}).xlsx'
    
    wb_s = openpyxl.load_workbook(f_sum, data_only=True)
    ws_s = wb_s.active
    h_s = [c for c in next(ws_s.iter_rows(values_only=True))]
    
    row_s = None
    for r in ws_s.iter_rows(values_only=True):
        if r[h_s.index('Store')] and str(r[h_s.index('Store')]) == '19762':
            row_s = r
            break
            
    if row_s:
        print(d, 'Orders:', row_s[h_s.index('Order Count')])
    else:
        print(d, 'STORE NOT FOUND!')
